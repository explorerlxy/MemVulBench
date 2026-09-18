from __future__ import annotations

import re
import tempfile
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'papers/MemVulBench/manuscript.md'
TEMPLATE = ROOT / 'papers/Template/《信息工程大学学报》论文模板.docx'
OUTPUT = ROOT / 'papers/MemVulBench/MemVulBench-信息工程大学学报版.docx'


def set_run_font(run, size=10.5, bold=False, east='宋体', west='Times New Roman'):
    run.font.name = west
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), east)


def add_text(paragraph, text, size=10.5, bold=False, east='宋体'):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, east)
    return run


def add_hyperlink(paragraph, label, url, size=9):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:ascii'), 'Times New Roman')
    rfonts.set(qn('w:hAnsi'), 'Times New Roman')
    rpr.append(rfonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = label
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|\\\(.*?\\\))')


def latex_plain(value):
    s = value.replace('\\(', '').replace('\\)', '')
    replacements = {
        r'\mathcal{V}': 'V', r'\mathrm{mem}': 'mem', r'\mathrm{known}': 'known',
        r'\mathrm{return}': 'return', r'\textit{类型}': '类型', r'\textit{方向}': '方向',
        r'\textit{文件}': '文件', r'\textit{行号}': '行号', r'\subseteq': '⊆',
        r'\leq': '≤', r'\preceq': '≼', r'\not\preceq': '⊀',
        r'\sum': 'Σ', r'\in': '∈', r'\mathbb{1}': '1', r'\Vert': '∥',
        r'\ldots': '…', r'\bmod': 'mod', r'\qquad': '    ', r'\,': ' ',
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    s = s.replace('{', '').replace('}', '')
    return s


def inline(paragraph, text, size=10.5, east='宋体'):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            add_text(paragraph, part[2:-2], size, True, '黑体' if east == '宋体' else east)
        elif part.startswith('`') and part.endswith('`'):
            add_text(paragraph, part[1:-1], size, False, east)
        elif part.startswith(r'\(') and part.endswith(r'\)'):
            add_text(paragraph, latex_plain(part), size, False, east)
        elif part.startswith('[') and '](' in part:
            label, url = part[1:-1].split('](', 1)
            add_hyperlink(paragraph, label, url, size)
        else:
            add_text(paragraph, part, size, False, east)


def set_columns(section, count):
    cols = section._sectPr.xpath('./w:cols')
    if cols:
        el = cols[0]
    else:
        el = OxmlElement('w:cols')
        section._sectPr.append(el)
    el.set(qn('w:num'), str(count))
    el.set(qn('w:space'), '425')


def section(doc, count):
    sec = doc.add_section(WD_SECTION_START.CONTINUOUS)
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(1.9)
    sec.left_margin, sec.right_margin = Cm(1.9), Cm(1.9)
    sec.header_distance, sec.footer_distance = Cm(1), Cm(1)
    set_columns(sec, count)
    return sec


def base_para(p, indent=True, after=3):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(21) if indent else Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.15
    pf.widow_control = True


def body(doc, text):
    p = doc.add_paragraph()
    base_para(p)
    inline(p, text)


def heading(doc, text, level):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
    p.paragraph_format.line_spacing = 1.1
    add_text(p, text, 14 if level == 1 else 10.5, level != 1, '宋体' if level == 1 else '黑体')


def caption(doc, text, kind):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == 'figure' else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = kind == 'table'
    inline(p, text, 8.5 if kind == 'figure' else 7.5, '宋体' if kind == 'figure' else '黑体')


def figure(doc, path):
    match = re.fullmatch(r'!\[[^\]]*\]\(([^)]+)\)', path)
    if match is None:
        raise ValueError(f'Invalid image reference: {path}')
    fp = Path(match.group(1))
    if not fp.is_absolute():
        fp = SOURCE.parent / fp
    if not fp.exists():
        raise FileNotFoundError(fp)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(fp), width=Cm(16))


def border(cell, edge, val='single', sz='6'):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcpr.append(borders)
    el = borders.find(qn('w:' + edge))
    if el is None:
        el = OxmlElement('w:' + edge)
        borders.append(el)
    el.set(qn('w:val'), val)
    el.set(qn('w:sz'), sz)
    el.set(qn('w:color'), '000000')


def table(doc, rows):
    cols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = ({4: [Cm(3.3), Cm(4.6), Cm(6.0), Cm(3.3)],
               5: [Cm(2.3), Cm(3.5), Cm(3.7), Cm(3.8), Cm(3.9)],
               3: [Cm(3.4), Cm(8.0), Cm(5.8)]})[cols]
    if cols == 4 and len(rows) > 10:
        widths = [Cm(3.6), Cm(4.0), Cm(5.5), Cm(4.1)]
    for j, w in enumerate(widths):
        t.columns[j].width = w
    for i, row in enumerate(rows):
        trpr = t.rows[i]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trpr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            trpr.append(repeat)
        for j, content in enumerate(row):
            c = t.cell(i, j)
            c.width = widths[j]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or (cols == 4 and len(rows) > 10 and j in (0, 1, 3))) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            inline(p, content, 8 if cols == 5 else 8.5, '黑体' if i == 0 else '宋体')
            if i == 0:
                for r in p.runs:
                    r.font.bold = True
            if i == 0:
                border(c, 'top', sz='10')
                border(c, 'bottom', sz='6')
            if i == len(rows) - 1:
                border(c, 'bottom', sz='10')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


FORMULAS = [
    '(类型, 方向, 文件, 行号, H₃)',
    'C = (F, S, U, T)',
    'B(U; S, E) ⊆ V_mem(U; E),\nn_known(U; S, E) = |B(U; S, E)| ≤ |V_mem(U; E)|',
    'Q(c) = Σ_(b∈C_P) 1{f_b ⊀ c} ·\n1{0 < t(f_b) − t(c) ≤ 730 days}',
    'H(x) = { return, if |x| = 0;\nh_(x₀ mod k)(x₁…x_(n−1)), if |x| > 0 }',
]


def equation(doc, index):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    add_text(p, FORMULAS[index], 9.5)
    add_text(p, f'  （{index+1}）', 9)


def renumber(text):
    def cross(m):
        return f'第{int(m.group(1))-1}{m.group(2) or ""}节'
    return re.sub(r'第([2-8])(\.\d+)?节', cross, text)


def parse_table(lines, start):
    rows = []
    j = start
    while j < len(lines) and lines[j].startswith('|'):
        cells = [x.strip() for x in lines[j].strip().strip('|').split('|')]
        if not all(re.fullmatch(r':?-+:?', x) for x in cells):
            rows.append(cells)
        j += 1
    return rows, j


def clear_template_footer_text(path: Path) -> None:
    """Clear sample footer content left in unreferenced template footer parts."""
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.docx', delete=False) as temp:
        temporary = Path(temp.name)
    try:
        with ZipFile(path) as source, ZipFile(temporary, 'w') as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if re.fullmatch(r'word/footer\d+\.xml', item.filename):
                    root = ElementTree.fromstring(content)
                    for node in root.iter(namespace + 't'):
                        node.text = ''
                    content = ElementTree.tostring(root, encoding='utf-8', xml_declaration=True)
                target.writestr(item, content)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def build():
    doc = Document(TEMPLATE)
    body_el = doc._element.body
    sectpr = body_el.sectPr
    for child in list(body_el):
        if child is not sectpr:
            body_el.remove(child)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    set_columns(sec, 1)
    for p in sec.footer.paragraphs:
        p.clear()
    for p in sec.first_page_footer.paragraphs:
        p.clear()
    for p in sec.header.paragraphs:
        p.clear()
    doc.core_properties.title = 'MemVulBench 面向C/C++内存漏洞模糊测试的真实多漏洞基准构建'
    doc.core_properties.subject = '信息工程大学学报格式论文稿'
    doc.core_properties.author = ''

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    title = lines[0][2:]
    english = lines[2].strip('*')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    add_text(p, title, 18, False, '宋体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_text(p, '作者与单位：待补充', 9, False, '楷体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    add_text(p, '摘要：', 9, True, '黑体')
    inline(p, lines[10], 9, '仿宋')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    inline(p, lines[12], 9, '仿宋')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    add_text(p, english, 14, True, '宋体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    add_text(p, 'Authors and affiliations to be supplied', 9, False, '宋体')
    p = doc.add_paragraph()
    add_text(p, 'Abstract: ', 9, True)
    inline(p, lines[16], 9)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    inline(p, lines[18], 9)
    section(doc, 2)
    i = next(i for i, line in enumerate(lines) if line == '## 1 引言')
    formula_index = 0
    while i < len(lines):
        raw = lines[i]
        if not raw.strip():
            i += 1
            continue
        if raw.startswith('## '):
            name = raw[3:].strip()
            if name.startswith('1 引言'):
                name = '引言'
            elif re.match(r'^[2-8] ', name):
                first, rest = name.split(' ', 1)
                name = f'{int(first)-1}  {rest}'
            heading(doc, name, 1)
        elif raw.startswith('### '):
            name = raw[4:].strip()
            name = re.sub(r'^([2-8])(\.\d+)', lambda m: str(int(m.group(1))-1) + m.group(2), name)
            heading(doc, name, 2)
        elif raw.startswith('!['):
            section(doc, 1)
            figure(doc, raw)
            i += 2  # blank line before the caption
            caption(doc, renumber(lines[i]), 'figure')
            section(doc, 2)
        elif re.match(r'^表\s*\d+[\s　]', raw):
            section(doc, 1)
            caption(doc, renumber(raw), 'table')
            i += 2  # blank line before table
            rows, j = parse_table(lines, i)
            table(doc, rows)
            section(doc, 2)
            i = j - 1
        elif raw == r'\[':
            while i < len(lines) and lines[i] != r'\]':
                i += 1
            equation(doc, formula_index)
            formula_index += 1
        elif raw.startswith('|'):
            rows, j = parse_table(lines, i)
            table(doc, rows)
            i = j - 1
        elif raw.startswith('[') and re.match(r'^\[\d+\]', raw):
            p = doc.add_paragraph()
            base_para(p, False, 3)
            p.paragraph_format.left_indent = Pt(20)
            p.paragraph_format.first_line_indent = Pt(-20)
            inline(p, raw, 9)
        elif raw.startswith('>'):
            # The source's internal revision memo is kept in Markdown, not in the submission manuscript.
            pass
        else:
            body(doc, renumber(raw))
        i += 1
    if formula_index != len(FORMULAS):
        raise RuntimeError(f'Expected {len(FORMULAS)} formulas, found {formula_index}')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    clear_template_footer_text(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
